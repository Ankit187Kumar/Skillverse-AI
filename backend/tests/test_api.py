import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.interview_service import target_question_count
from app.services.llm import LLMError, extract_json, generate_json
from app.schemas import ResumeParsed

from .conftest import register

CONFIG = {"interview_type": "Technical", "role": "Backend Developer", "skills": ["Python", "SQL"],
          "experience_level": "Intermediate", "difficulty": "Medium", "duration_minutes": 9, "language": "English"}


# ---------- auth ----------
def test_register_login_me_logout(client):
    u = register(client)
    assert u["email"] == "asha@example.com" and u["streak_count"] == 1
    assert client.get("/api/auth/me").status_code == 200
    assert client.post("/api/auth/logout").status_code == 204
    assert client.get("/api/auth/me").status_code == 401
    r = client.post("/api/auth/login", json={"email": "ASHA@example.com", "password": "secret123"})
    assert r.status_code == 200
    assert client.get("/api/auth/me").json()["full_name"] == "Asha Rao"


def test_register_rejects_duplicates_and_weak_passwords(client):
    register(client)
    r = client.post("/api/auth/register", json={"email": "asha@example.com", "password": "secret123", "full_name": "X"})
    assert r.status_code == 409
    r = client.post("/api/auth/register", json={"email": "b@example.com", "password": "onlyletters", "full_name": "B"})
    assert r.status_code == 422 and "Password" in r.json()["detail"]


def test_wrong_password(client):
    register(client)
    client.post("/api/auth/logout")
    r = client.post("/api/auth/login", json={"email": "asha@example.com", "password": "nope12345"})
    assert r.status_code == 401


def test_password_is_hashed(client):
    from app.database import SessionLocal
    from app.models import User
    register(client)
    with SessionLocal() as db:
        u = db.query(User).one()
        assert u.password_hash != "secret123" and u.password_hash.startswith("$2")


def test_csrf_header_required_with_cookie(client):
    register(client)
    with TestClient(app, cookies=client.cookies) as raw:  # no custom header
        r = raw.patch("/api/users/me", json={"headline": "hi"})
        assert r.status_code == 403


def test_login_rate_limited(client):
    register(client)
    codes = [client.post("/api/auth/login", json={"email": "asha@example.com", "password": "bad12345"}).status_code
             for _ in range(12)]
    assert 429 in codes


# ---------- profile ----------
def test_profile_update_and_export(client):
    register(client)
    r = client.patch("/api/users/me", json={"skills": ["Python", "python", " SQL "], "target_roles": ["SDE"],
                                            "experience_level": "Advanced", "learning_goals": ["Crack system design"]})
    assert r.status_code == 200
    assert r.json()["skills"] == ["Python", "SQL"] and r.json()["experience_level"] == "Advanced"
    export = client.get("/api/users/me/export").json()
    assert export["profile"]["target_roles"] == ["SDE"]


def test_change_password(client):
    register(client)
    assert client.post("/api/users/me/password", json={"current_password": "wrong", "new_password": "newpass99"}).status_code == 400
    assert client.post("/api/users/me/password", json={"current_password": "secret123", "new_password": "newpass99"}).status_code == 204
    client.post("/api/auth/logout")
    assert client.post("/api/auth/login", json={"email": "asha@example.com", "password": "newpass99"}).status_code == 200


# ---------- interviews ----------
def test_question_count_scales_with_duration():
    assert target_question_count(5) == 3
    assert target_question_count(30) == 10
    assert target_question_count(60) == 12


def test_full_interview_flow(client, llm):
    register(client)
    r = client.post("/api/interviews", json=CONFIG)
    assert r.status_code == 201, r.text
    iv = r.json()
    assert iv["target_questions"] == 3 and iv["questions_asked"] == 1 and iv["turns"][0]["role"] == "interviewer"

    for i in range(3):
        r = client.post(f"/api/interviews/{iv['id']}/answer", json={"answer": f"My answer {i}", "elapsed_seconds": 60 * (i + 1)})
        assert r.status_code == 200, r.text
    body = r.json()
    assert body["done"] is True and body["interview"]["questions_asked"] == 3

    r = client.post(f"/api/interviews/{iv['id']}/end", json={"elapsed_seconds": 200})
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["status"] == "ready"
    # weighted: 80*.25 + 70*.25 + 90*.2 + 60*.2 + 50*.1 = 72.5
    assert rep["overall_score"] == 72.5
    assert "confidence" in rep["criteria"] and "camera" in rep["criteria"]["confidence"]

    assert client.post(f"/api/interviews/{iv['id']}/answer", json={"answer": "late"}).status_code == 409
    listing = client.get("/api/interviews").json()
    assert listing[0]["overall_score"] == 72.5

    dash = client.get("/api/dashboard").json()
    assert dash["stats"]["interviews_completed"] == 1
    assert dash["stats"]["average_score"] == 72.5
    assert dash["stats"]["practice_minutes"] == 3
    assert any("Docker" in r["title"] for r in dash["recommendations"])
    assert any(a["id"] == "first_step" and a["earned"] for a in dash["achievements"])


def test_ending_early_drops_unanswered_question(client):
    register(client)
    iv = client.post("/api/interviews", json=CONFIG).json()
    client.post(f"/api/interviews/{iv['id']}/answer", json={"answer": "First answer"})
    rep = client.post(f"/api/interviews/{iv['id']}/end").json()
    assert rep["status"] == "ready"
    turns = client.get(f"/api/interviews/{iv['id']}").json()["turns"]
    assert turns[-1]["role"] == "candidate"


def test_end_without_answers_gives_clear_error(client):
    register(client)
    iv = client.post("/api/interviews", json=CONFIG).json()
    rep = client.post(f"/api/interviews/{iv['id']}/end").json()
    assert rep["status"] == "failed" and "No answers" in rep["error"]


def test_llm_failure_keeps_answer_and_can_retry(client, llm):
    register(client)
    iv = client.post("/api/interviews", json=CONFIG).json()
    llm.fail = True
    r = client.post(f"/api/interviews/{iv['id']}/answer", json={"answer": "Saved even if model fails"})
    assert r.status_code == 503 and r.json()["code"] == "ai_unavailable"
    turns = client.get(f"/api/interviews/{iv['id']}").json()["turns"]
    assert turns[-1]["content"] == "Saved even if model fails"
    llm.fail = False
    r = client.post(f"/api/interviews/{iv['id']}/next")
    assert r.status_code == 200 and r.json()["turns"][-1]["role"] == "interviewer"


def test_start_fails_cleanly_when_model_down(client, llm):
    register(client)
    llm.fail = True
    r = client.post("/api/interviews", json=CONFIG)
    assert r.status_code == 503
    assert client.get("/api/interviews").json() == []


def test_users_cannot_see_each_others_data(client):
    register(client)
    iv = client.post("/api/interviews", json=CONFIG).json()
    client.post("/api/auth/logout")
    register(client, email="ravi@example.com", name="Ravi")
    assert client.get(f"/api/interviews/{iv['id']}").status_code == 404
    assert client.post(f"/api/interviews/{iv['id']}/answer", json={"answer": "x"}).status_code == 404
    assert client.delete(f"/api/interviews/{iv['id']}").status_code == 404
    assert client.get("/api/interviews").json() == []


def test_invalid_config_rejected(client):
    register(client)
    r = client.post("/api/interviews", json={**CONFIG, "difficulty": "Impossible"})
    assert r.status_code == 422
    r = client.post("/api/interviews", json={**CONFIG, "duration_minutes": 500})
    assert r.status_code == 422


# ---------- resumes ----------
def _docx_bytes() -> bytes:
    d = Document()
    d.add_paragraph("Asha Rao — Backend Engineer")
    d.add_paragraph("Skills: Python, FastAPI, PostgreSQL, Docker. Built REST APIs serving 1M requests/day.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_resume_upload_parse_and_use(client, llm):
    register(client)
    r = client.post("/api/resumes", files={"file": ("cv.docx", _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 201, r.text
    res = r.json()
    assert res["parsed"]["skills"] == ["Python", "FastAPI"] and "Backend Engineer" in res["text_preview"]
    iv = client.post("/api/interviews", json={**CONFIG, "resume_id": res["id"]}).json()
    assert iv["config"]["resume_id"] == res["id"]
    assert "Resume summary" in llm.calls[-1][0]["content"]


def test_resume_rejects_spoofed_and_large_files(client):
    register(client)
    r = client.post("/api/resumes", files={"file": ("cv.pdf", b"not really a pdf" * 10, "application/pdf")})
    assert r.status_code == 422
    r = client.post("/api/resumes", files={"file": ("cv.exe", b"MZ....", "application/octet-stream")})
    assert r.status_code == 422
    big = b"%PDF-" + b"0" * (6 * 1024 * 1024)
    r = client.post("/api/resumes", files={"file": ("cv.pdf", big, "application/pdf")})
    assert r.status_code == 413


def test_resume_of_other_user_cannot_be_used(client):
    register(client)
    res = client.post("/api/resumes", files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")}).json()
    client.post("/api/auth/logout")
    register(client, email="ravi@example.com", name="Ravi")
    r = client.post("/api/interviews", json={**CONFIG, "resume_id": res["id"]})
    assert r.status_code == 400
    assert client.delete(f"/api/resumes/{res['id']}").status_code == 404


# ---------- account deletion ----------
def test_delete_account_removes_everything(client):
    from app.database import SessionLocal
    from app.models import Activity, Interview, InterviewTurn, Resume
    register(client)
    client.post("/api/resumes", files={"file": ("cv.docx", _docx_bytes(), "application/octet-stream")})
    iv = client.post("/api/interviews", json=CONFIG).json()
    client.post(f"/api/interviews/{iv['id']}/answer", json={"answer": "a"})
    assert client.request("DELETE", "/api/users/me", json={"password": "wrong"}).status_code == 400
    assert client.request("DELETE", "/api/users/me", json={"password": "secret123"}).status_code == 204
    assert client.get("/api/auth/me").status_code == 401
    with SessionLocal() as db:
        for model in (Resume, Interview, InterviewTurn, Activity):
            assert db.query(model).count() == 0


# ---------- structured output ----------
def test_extract_json_handles_fences_and_chatter():
    assert extract_json('Sure!\n```json\n{"a": 1}\n```') == {"a": 1}
    assert extract_json('blah {"b": 2} trailing') == {"b": 2}
    with pytest.raises(ValueError):
        extract_json("no json here")


def test_generate_json_retries_then_fails():
    class Bad:
        model = "x"
        def __init__(self): self.n = 0
        def chat(self, messages, **kw):
            self.n += 1
            return '{"skills": "not a list"}' if self.n == 1 else '{"skills": ["Go"]}'
    out = generate_json(Bad(), [{"role": "user", "content": "x"}], ResumeParsed)
    assert out.skills == ["Go"]

    class AlwaysBad(Bad):
        def chat(self, messages, **kw):
            return "nope"
    with pytest.raises(LLMError):
        generate_json(AlwaysBad(), [{"role": "user", "content": "x"}], ResumeParsed)


def test_health_and_ai_status(client):
    assert client.get("/api/health").json()["status"] == "ok"
    assert client.get("/api/system/ai-status").status_code == 401
    register(client)
    assert client.get("/api/system/ai-status").json()["model_available"] is True


def test_timestamps_are_utc_aware(client):
    u = register(client)
    assert u["created_at"].endswith(("+00:00", "Z"))
