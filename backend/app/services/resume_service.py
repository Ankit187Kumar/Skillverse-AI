"""Resume upload validation and text extraction (PDF and DOCX)."""

from __future__ import annotations

import io
import zipfile

from pypdf import PdfReader
from pypdf.errors import PdfReadError

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"


class ResumeError(ValueError):
    pass


def detect_type(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf") and data.startswith(PDF_MAGIC):
        return "pdf"
    if name.endswith(".docx") and data.startswith(ZIP_MAGIC):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    raise ResumeError("Only PDF and DOCX files are supported here, and the file contents must match the extension.")


def _guard_zip_bomb(data: bytes, limit: int = 50 * 1024 * 1024) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > limit:
            raise ResumeError("This DOCX file expands to an unusually large size and was rejected.")


def extract_text(filename: str, data: bytes, *, max_pages: int = 15, max_chars: int = 30000,
                 label: str = "resumes") -> str:
    kind = detect_type(filename, data)
    try:
        if kind == "pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeError("Password-protected PDFs are not supported.")
            pages = reader.pages[:max_pages]
            text = "\n".join((p.extract_text() or "") for p in pages)
        else:
            _guard_zip_bomb(data)
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
    except ResumeError:
        raise
    except (PdfReadError, KeyError, ValueError, zipfile.BadZipFile) as e:
        raise ResumeError("The file could not be read. It may be damaged.") from e

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 40:
        raise ResumeError(f"Very little text was found. Scanned (image-only) {label} are not supported yet.")
    return text[:max_chars]
